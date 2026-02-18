"""
Generate architecture.pptx for the Factory Inventory Management System.
Run from the project root: python scripts/gen_architecture_ppt.py
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.enum.shapes import MSO_SHAPE_TYPE
import copy

# ── Design tokens ──────────────────────────────────────────────────────────────
COLOR_DARK   = RGBColor(0x0F, 0x17, 0x2A)   # #0f172a – headings
COLOR_MID    = RGBColor(0x64, 0x74, 0x8B)   # #64748b – body text
COLOR_LIGHT  = RGBColor(0xE2, 0xE8, 0xF0)   # #e2e8f0 – rule lines
COLOR_ACCENT = RGBColor(0x3B, 0x82, 0xF6)   # #3b82f6 – blue accent
COLOR_WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_GREEN  = RGBColor(0x10, 0xB9, 0x81)   # #10b981
COLOR_AMBER  = RGBColor(0xF5, 0x9E, 0x0B)   # #f59e0b

W = 13.33   # slide width  (inches)
H = 7.5     # slide height (inches)

# ── Helpers ────────────────────────────────────────────────────────────────────

def _blank_slide(prs):
    blank_layout = prs.slide_layouts[6]   # index 6 = totally blank
    return prs.slides.add_slide(blank_layout)


def _set_font(run, size_pt, bold=False, italic=False, color=None):
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color if color else COLOR_DARK


def _add_rect(slide, left, top, width, height, fill_color, border=False):
    from pptx.util import Inches
    shape = slide.shapes.add_shape(
        1,  # MSO_AUTO_SHAPE_TYPE.RECTANGLE
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if not border:
        shape.line.fill.background()
    return shape


def _add_textbox(slide, text, left, top, width, height,
                 size=14, bold=False, italic=False,
                 color=None, align=PP_ALIGN.LEFT, wrap=True):
    txBox = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    tf = txBox.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_font(run, size, bold=bold, italic=italic, color=color or COLOR_DARK)
    return txBox


def _add_bullet_textbox(slide, items, left, top, width, height,
                        size=13, color=None, indent_char="  •  "):
    txBox = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.space_before = Pt(3)
        run = p.add_run()
        run.text = f"{indent_char}{item}"
        _set_font(run, size, color=color or COLOR_MID)
    return txBox


# ── Slide builders ──────────────────────────────────────────────────────────────

def add_title_slide(prs):
    slide = _blank_slide(prs)

    # Dark full-bleed background
    _add_rect(slide, 0, 0, W, H, COLOR_DARK)

    # Accent bar (blue strip, left edge)
    _add_rect(slide, 0, 0, 0.18, H, COLOR_ACCENT)

    # Title
    _add_textbox(slide, "Factory Inventory Management System",
                 0.55, 2.2, 11.5, 1.4,
                 size=36, bold=True, color=COLOR_WHITE, align=PP_ALIGN.LEFT)

    # Subtitle
    _add_textbox(slide, "Application Architecture Overview",
                 0.55, 3.7, 10, 0.8,
                 size=22, bold=False, color=RGBColor(0xA0, 0xAE, 0xC0),
                 align=PP_ALIGN.LEFT)

    # Footer tag
    _add_textbox(slide, "Vue 3  +  FastAPI  +  In-Memory Data  ·  Full-Stack Demo",
                 0.55, 6.5, 12, 0.5,
                 size=12, italic=True,
                 color=RGBColor(0x64, 0x74, 0x8B), align=PP_ALIGN.LEFT)


def _content_header(slide, title, subtitle=None):
    """Shared header: accent bar + title + rule line."""
    _add_rect(slide, 0, 0, W, 0.08, COLOR_ACCENT)   # thin top bar
    _add_textbox(slide, title, 0.5, 0.2, W - 1, 0.75,
                 size=24, bold=True, color=COLOR_DARK)
    if subtitle:
        _add_textbox(slide, subtitle, 0.5, 0.9, W - 1, 0.4,
                     size=13, italic=True, color=COLOR_MID)
    # rule line
    _add_rect(slide, 0.5, 1.0, W - 1, 0.02, COLOR_LIGHT)


def add_content_slide(prs, title, bullets, subtitle=None):
    slide = _blank_slide(prs)
    _content_header(slide, title, subtitle)
    _add_bullet_textbox(slide, bullets, 0.5, 1.25, W - 1, H - 1.5, size=14)
    return slide


def add_two_column_slide(prs, title,
                         left_heading, left_items,
                         right_heading, right_items,
                         subtitle=None):
    slide = _blank_slide(prs)
    _content_header(slide, title, subtitle)

    col_w = (W - 1.5) / 2
    top = 1.3

    # Left column
    _add_textbox(slide, left_heading, 0.5, top, col_w, 0.45,
                 size=14, bold=True, color=COLOR_ACCENT)
    _add_rect(slide, 0.5, top + 0.45, col_w, 0.02, COLOR_LIGHT)
    _add_bullet_textbox(slide, left_items, 0.5, top + 0.55, col_w, H - top - 1,
                        size=13)

    # Divider
    _add_rect(slide, W / 2, top, 0.02, H - top - 0.3, COLOR_LIGHT)

    # Right column
    right_left = W / 2 + 0.2
    _add_textbox(slide, right_heading, right_left, top, col_w, 0.45,
                 size=14, bold=True, color=COLOR_ACCENT)
    _add_rect(slide, right_left, top + 0.45, col_w, 0.02, COLOR_LIGHT)
    _add_bullet_textbox(slide, right_items, right_left, top + 0.55,
                        col_w, H - top - 1, size=13)
    return slide


# ── Build deck ──────────────────────────────────────────────────────────────────

def build(output_path="architecture.pptx"):
    prs = Presentation()
    prs.slide_width  = Inches(W)
    prs.slide_height = Inches(H)

    # ── Slide 1 – Title ────────────────────────────────────────────────────────
    add_title_slide(prs)

    # ── Slide 2 – System Overview ──────────────────────────────────────────────
    add_content_slide(prs,
        title="System Architecture Overview",
        bullets=[
            "Full-stack web application for factory inventory management",
            "Three-tier architecture:  Frontend  ·  Backend API  ·  Data Layer",
            "Frontend:  Vue 3 (port 3000) served by Vite dev server",
            "Backend:  Python FastAPI (port 8001) with Uvicorn ASGI server",
            "Data:  JSON flat files loaded into memory at startup — no database",
            "Communication:  REST API over HTTP via Axios client",
        ]
    )

    # ── Slide 3 – Frontend ─────────────────────────────────────────────────────
    add_two_column_slide(prs,
        title="Frontend — Vue 3 + Vite",
        left_heading="Stack",
        left_items=[
            "Vue 3 Composition API (<script setup>)",
            "Vite build tool & dev server",
            "Vue Router for client-side navigation",
            "Axios for HTTP requests",
            "Scoped CSS — no external UI library",
            "Custom SVG charts (no chart library)",
        ],
        right_heading="Directory Structure",
        right_items=[
            "client/src/views/      — 7 page-level components",
            "client/src/components/ — Reusable UI components",
            "client/src/composables/— Shared reactive logic",
            "client/src/api.js      — Centralised API client",
            "client/src/App.vue     — Root layout + sidebar",
            "client/src/main.js     — App entry point",
        ]
    )

    # ── Slide 4 – Backend ──────────────────────────────────────────────────────
    add_two_column_slide(prs,
        title="Backend — Python FastAPI",
        left_heading="Stack",
        left_items=[
            "Python 3.11+ with FastAPI framework",
            "Uvicorn ASGI server",
            "Pydantic v2 for request/response validation",
            "CORS middleware (allows localhost:3000)",
            "In-memory data store — no ORM, no database",
            "uv for dependency management",
        ],
        right_heading="Directory Structure",
        right_items=[
            "server/main.py         — FastAPI app + all route handlers",
            "server/mock_data.py    — Data loading + filtering logic",
            "server/data/*.json     — Raw data files",
            "server/requirements.txt— Python dependencies",
            "All data loaded once at startup into module-level variables",
        ]
    )

    # ── Slide 5 – Data Layer ───────────────────────────────────────────────────
    add_content_slide(prs,
        title="Data Layer — JSON + In-Memory",
        bullets=[
            "No database — all data lives in server/data/ as JSON flat files",
            "Files:  orders.json · inventory.json · backlog.json · demand.json · spending.json",
            "mock_data.py reads all files once at startup into Python lists/dicts",
            "Filtering is done in-memory using Python list comprehensions",
            "Pydantic models validate every API response before it leaves the server",
            "Data mutations (e.g. creating a PO) update in-memory state only — resets on restart",
        ]
    )

    # ── Slide 6 – Data Flow ────────────────────────────────────────────────────
    add_content_slide(prs,
        title="End-to-End Data Flow",
        bullets=[
            "1 →  User changes a filter in the FilterBar (Vue reactive ref updates)",
            "2 →  useFilters composable broadcasts the change to all views via shared state",
            "3 →  Each view calls its loadData() function with current filter values",
            "4 →  client/src/api.js builds the query string and fires an Axios GET request",
            "5 →  FastAPI receives the request, extracts query params, passes to mock_data.py",
            "6 →  mock_data.py filters in-memory lists and returns matching records",
            "7 →  FastAPI validates the result through Pydantic models and serialises to JSON",
            "8 →  Axios resolves the promise; Vue stores raw data in ref() variables",
            "9 →  Computed properties derive display values; Vue re-renders reactively",
        ]
    )

    # ── Slide 7 – Filter System ────────────────────────────────────────────────
    add_content_slide(prs,
        title="Global Filter System",
        bullets=[
            "Four global filters apply across all pages simultaneously",
            "Time Period  — month selector (Jan–Dec or All Months)  →  'month' query param",
            "Location     — warehouse selector (San Francisco / London / Tokyo)  →  'warehouse'",
            "Category     — product category  →  'category' query param",
            "Order Status — Delivered / Shipped / Processing / Backordered  →  'status'",
            "Filters stored in useFilters composable — single source of truth across all views",
            "Inventory endpoints ignore 'month' (no time dimension on stock levels)",
            "Revenue goal auto-adjusts:  $800K/month  or  $9.6M for the full year",
        ]
    )

    # ── Slide 8 – API Endpoints ────────────────────────────────────────────────
    add_content_slide(prs,
        title="REST API Endpoints",
        subtitle="Base URL:  http://localhost:8001",
        bullets=[
            "GET /api/inventory              — Stock levels          filters: warehouse, category",
            "GET /api/inventory/{id}         — Single item detail",
            "GET /api/orders                 — Order list            filters: warehouse, category, status, month",
            "GET /api/orders/{id}            — Single order detail",
            "GET /api/dashboard/summary      — KPI aggregates        all 4 filters",
            "GET /api/demand                 — Demand forecast data  no filters",
            "GET /api/backlog                — Shortage/backlog items no filters",
            "GET /api/spending/summary       — Spend overview        no filters",
            "GET /api/spending/monthly       — Monthly spend breakdown",
            "GET /api/spending/categories    — Spend by category",
            "GET /api/spending/transactions  — Transaction list",
        ]
    )

    # ── Slide 9 – Pages / Views ────────────────────────────────────────────────
    add_two_column_slide(prs,
        title="Application Pages",
        left_heading="Routes",
        left_items=[
            "Overview (Dashboard)  —  /",
            "Inventory             —  /inventory",
            "Orders                —  /orders",
            "Finance               —  /spending",
            "Demand Forecast       —  /demand",
            "Restocking            —  /restocking",
            "Reports               —  /reports",
        ],
        right_heading="Dashboard Sections",
        right_items=[
            "KPI cards (turnover rate, fill rate, revenue, processing time)",
            "Order health donut chart + metrics",
            "Inventory value by category bar chart",
            "Order trend line chart (monthly)",
            "Inventory shortages table with Create PO actions",
            "Top products by revenue table",
        ]
    )

    # ── Slide 10 – Key Patterns ────────────────────────────────────────────────
    add_content_slide(prs,
        title="Key Design Patterns",
        bullets=[
            "Composable pattern  —  useFilters · useAuth · useI18n share state across components",
            "Ref / Computed split  —  raw API data in ref(), derived display values in computed()",
            "Centralised API client  —  all HTTP calls go through api.js, never directly from components",
            "Filter propagation  —  watch([...filters], loadData) triggers reload on any filter change",
            "Unique v-for keys  —  always use sku · order_id · month  (never array index)",
            "Date validation  —  every new Date(str) call guarded with isNaN(date.getTime())",
            "Modal pattern  —  <Teleport to='body'> + <Transition> for all overlay dialogs",
            "No external UI library  —  custom scoped CSS with slate/gray design tokens",
        ]
    )

    prs.save(output_path)
    slide_count = len(prs.slides)
    print(f"Saved: {output_path}  ({slide_count} slides)")
    return output_path


if __name__ == "__main__":
    build("architecture.pptx")
